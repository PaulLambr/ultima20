import json
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from zipfile import BadZipFile

SNAPSHOT_FOLDER = Path.home() / "Ledger_History"
MASTER_EXCEL_FILE = SNAPSHOT_FOLDER / "monthly_ledger.xlsx"

# Delete and reset corrupted Excel file
if MASTER_EXCEL_FILE.exists():
    try:
        wb = load_workbook(MASTER_EXCEL_FILE)
    except BadZipFile:
        print("⚠️ Corrupted Excel file detected. Deleting and starting fresh.")
        MASTER_EXCEL_FILE.unlink()

# Load processed filenames
if MASTER_EXCEL_FILE.exists():
    try:
        wb = load_workbook(MASTER_EXCEL_FILE)
        if "IngestedFiles" in wb.sheetnames:
            processed = pd.read_excel(MASTER_EXCEL_FILE, sheet_name="IngestedFiles")["Filename"].tolist()
        else:
            processed = []
    except:
        processed = []
else:
    processed = []

# Find new snapshot files
snapshot_files = sorted(SNAPSHOT_FOLDER.glob("snapshot-*.json"), key=lambda f: f.stat().st_mtime)
new_snapshots = [f for f in snapshot_files if f.name not in processed]

if not new_snapshots:
    print("⚠️ No new snapshot files to process.")
    exit()

processed_updates = []

with pd.ExcelWriter(MASTER_EXCEL_FILE, engine="openpyxl", mode="a" if MASTER_EXCEL_FILE.exists() else "w") as writer:
    for file in new_snapshots:
        with open(file) as f:
            data = json.load(f)

        timestamp = data["timestamp"][:10]
        net_total = data["netTotal"]
        asset_total = data["assetTotal"]
        bill_total = data["billTotal"]
        assets_df = pd.DataFrame(data["assets"])
        ledger_df = pd.DataFrame(data["ledger"])

        print(f"📁 Processing snapshot: {file.name}")
        print(f"📊 Timestamp: {timestamp}, Net: {net_total}")

        # --- Handle Summary tab ---
        summary_row = pd.DataFrame([{
            "Date": timestamp,
            "Net Total": net_total,
            "Asset Total": asset_total,
            "Bill Total": bill_total
        }])

        try:
            existing_summary = pd.read_excel(MASTER_EXCEL_FILE, sheet_name="Summary")
            full_summary = pd.concat([existing_summary, summary_row], ignore_index=True)
        except:
            full_summary = summary_row

        if "Summary" in writer.book.sheetnames:
            if len(writer.book.sheetnames) > 1:
                writer.book.remove(writer.book["Summary"])
            else:
                print("⚠️ Cannot remove 'Summary' as it's the only sheet. Skipping...")
                continue

        full_summary.to_excel(writer, sheet_name="Summary", index=False)

        # --- Write snapshot tab ---
        tab_name = timestamp
        counter = 1
        while tab_name in writer.book.sheetnames:
            tab_name = f"{timestamp}-{counter}"
            counter += 1

        sheet = writer.book.create_sheet(title=tab_name)

        # Assets
        for r_idx, row in enumerate(dataframe_to_rows(assets_df, index=False, header=True), start=1):
            for c_idx, value in enumerate(row, start=1):
                sheet.cell(row=r_idx, column=c_idx, value=value)

        # Ledger (after gap)
        ledger_start = len(assets_df) + 3
        for r_idx, row in enumerate(dataframe_to_rows(ledger_df, index=False, header=True), start=ledger_start):
            for c_idx, value in enumerate(row, start=1):
                sheet.cell(row=r_idx, column=c_idx, value=value)

        # Record this file
        processed_updates.append({
            "Filename": file.name,
            "Processed On": datetime.now().isoformat()
        })

    # --- IngestedFiles tab update ---
    updates_df = pd.DataFrame(processed_updates)
    try:
        previous = pd.read_excel(MASTER_EXCEL_FILE, sheet_name="IngestedFiles")
        combined = pd.concat([previous, updates_df], ignore_index=True)
        if "IngestedFiles" in writer.book.sheetnames:
            writer.book.remove(writer.book["IngestedFiles"])
    except:
        combined = updates_df

    combined.to_excel(writer, sheet_name="IngestedFiles", index=False)

print(f"✅ {len(new_snapshots)} file(s) processed and saved to {MASTER_EXCEL_FILE}")
